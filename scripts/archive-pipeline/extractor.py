# extractor.py
import re
from pathlib import Path


def extract_text(path: Path) -> str:
    ext = path.suffix.lstrip(".").lower()
    _extractors = {
        "docx": _docx,
        "pdf": _pdf,
        "pptx": _pptx,
        "txt": _txt,
        "xlsx": _xlsx,
    }
    fn = _extractors.get(ext)
    if not fn:
        return ""
    try:
        return clean_text(fn(path))
    except Exception:
        return ""


def clean_text(text: str) -> str:
    # Strip base64 image embeds (markdown and raw URI forms)
    text = re.sub(
        r'!\[.*?\]\(data:image/[^;]+;base64,[A-Za-z0-9+/=]+\)',
        "",
        text,
    )
    text = re.sub(r'data:image/[^;]+;base64,[A-Za-z0-9+/=]+', "", text)
    # Strip HTML tags
    text = re.sub(r'<[^>]+>', " ", text)
    # Collapse excessive whitespace
    text = re.sub(r'\n{3,}', "\n\n", text)
    text = re.sub(r' {2,}', " ", text)
    return text.strip()


def _docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _pdf(path: Path) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n".join(pages)


def _pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
    return "\n".join(texts)


def _txt(path: Path) -> str:
    return path.read_text(errors="ignore")


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append("\t".join(cells))
    wb.close()
    return "\n".join(rows)
